import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook

from generate import generate


class GenerationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "input.csv"
        self.template = self.root / "template.txt"
        self.output = self.root / "result"
        self.template.write_text("Service update\n\nHello {{organization}}\n\n{{record_id}}", encoding="utf-8")

    def test_content_and_metadata(self):
        self.source.write_text('record_id,organization\nDEMO-001,"Example & Café <A>"\n', encoding="utf-8")
        self.assertEqual(generate(self.source, self.template, self.output), 1)
        document = Document(self.output / "DEMO-001.docx")
        self.assertEqual(document.paragraphs[1].text, "Hello Example & Café <A>")
        self.assertFalse(document.core_properties.author)
        self.assertFalse(document.core_properties.last_modified_by)
        with ZipFile(self.output / "DEMO-001.docx") as archive:
            core = archive.read("docProps/core.xml").decode()
            self.assertNotIn("2013-", core)
            self.assertNotIn("docProps/thumbnail.jpeg", archive.namelist())

    def test_invalid_rows_leave_no_output(self):
        for rows in ("A,Example\nA,Example B\n", "A,Example\na,Example B\n", "../A,Example\n", "A,\n", "A,Example,extra\n"):
            with self.subTest(rows=rows):
                self.source.write_text("record_id,organization\n" + rows, encoding="utf-8")
                with self.assertRaises(ValueError):
                    generate(self.source, self.template, self.output)
                self.assertFalse(self.output.exists())

    def test_existing_output_is_untouched(self):
        self.output.mkdir()
        sentinel = self.output / "keep.txt"
        sentinel.write_text("keep")
        with self.assertRaises(ValueError):
            generate(self.source, self.template, self.output)
        self.assertEqual(sentinel.read_text(), "keep")

    def test_bad_placeholder_and_duplicate_headers(self):
        self.source.write_text("record_id,organization\nA,Example\n")
        self.template.write_text("Hello {{organization")
        with self.assertRaises(ValueError):
            generate(self.source, self.template, self.output)
        self.source.write_text("record_id,record_id\nA,B\n")
        with self.assertRaises(ValueError):
            generate(self.source, self.template, self.output)

    def test_xlsx_and_formula_rejection(self):
        source = self.root / "input.xlsx"
        workbook = Workbook()
        workbook.active.append(["record_id", "organization"])
        workbook.active.append(["DEMO-002", "Example Workshop"])
        workbook.save(source)
        self.assertEqual(generate(source, self.template, self.output), 1)
        workbook.active["B2"] = '=CONCAT("Example", " Workshop")'
        workbook.save(source)
        with self.assertRaisesRegex(ValueError, "Formula"):
            generate(source, self.template, self.root / "formula-output")
        self.assertFalse((self.root / "formula-output").exists())


if __name__ == "__main__":
    unittest.main()
