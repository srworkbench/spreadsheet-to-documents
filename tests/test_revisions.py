import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

import revisions
from review_page import render_page


class RevisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "records.csv"
        self.template = self.root / "template.txt"
        self.baseline = self.root / "v1"
        self.report = self.root / "review"
        self.next = self.root / "v2"
        self.source.write_text("record_id,date,note\nA,2030-06-15,internal\nB,2030-07-01,draft\n")
        self.template.write_text("Service renewal\n\nScheduled for {{date}}.")
        revisions.build(self.source, self.template, self.baseline)

    def plan(self):
        return revisions.review(self.baseline, self.source, self.template, self.report)

    def test_one_used_cell_rebuilds_one_document_unused_cell_reuses_bytes(self):
        original = {p.name: p.read_bytes() for p in self.baseline.iterdir()}
        self.source.write_text("record_id,date,note\nB,2030-07-01,changed\nA,2030-06-22,internal\n")
        plan = self.plan()
        self.assertEqual(plan["counts"], {"changed": 1, "unchanged": 1})
        self.assertEqual(plan["entries"][0]["fields"][0]["field"], "date")
        self.assertIn("rendered text is identical", (self.report / "index.html").read_text())
        with patch("revisions.write_document", wraps=revisions.write_document) as write:
            revisions.apply(self.report / "plan.json", self.next)
        self.assertEqual(write.call_count, 1)
        self.assertEqual((self.next / "B.docx").read_bytes(), original["B.docx"])
        self.assertIn("2030-06-22", Document(self.next / "A.docx").paragraphs[1].text)
        self.assertEqual({p.name:p.read_bytes() for p in self.baseline.iterdir()}, original)

    def test_template_change_rebuilds_both_and_then_can_be_reused(self):
        self.template.write_text("Service renewal\n\nYour new date is {{date}}.")
        plan = self.plan()
        self.assertTrue(plan["template_changed"])
        self.assertEqual(plan["counts"], {"changed": 2})
        revisions.apply(self.report / "plan.json", self.next)
        followup, _, _ = revisions.make_plan(self.next, self.source, self.template)
        self.assertEqual(followup["counts"], {"unchanged": 2})

    def test_hand_edited_word_file_blocks_even_if_its_record_was_removed(self):
        doc_path = self.baseline / "A.docx"
        doc = Document(doc_path)
        doc.add_paragraph("An invented manual note that must survive.")
        doc.save(doc_path)
        edited = doc_path.read_bytes()
        self.source.write_text("record_id,date,note\nB,2030-07-01,draft\n")
        plan = self.plan()
        self.assertTrue(plan["blocked"])
        self.assertEqual(plan["entries"][0]["action"], "removed")
        with self.assertRaisesRegex(ValueError, "baseline document"):
            revisions.apply(self.report / "plan.json", self.next)
        self.assertFalse(self.next.exists())
        self.assertEqual(doc_path.read_bytes(), edited)

    def test_added_and_removed_rows_leave_old_revision_untouched(self):
        self.source.write_text("record_id,date,note\nB,2030-07-01,draft\nC,2030-08-01,new\n")
        self.assertEqual(self.plan()["counts"], {"removed":1,"unchanged":1,"added":1})
        revisions.apply(self.report / "plan.json", self.next)
        self.assertEqual({p.name for p in self.next.glob("*.docx")}, {"B.docx","C.docx"})
        self.assertTrue((self.baseline / "A.docx").exists())

    def test_every_review_input_is_bound_and_stale_plans_do_not_write(self):
        for target in ("source", "template", "doc", "manifest", "plan"):
            with self.subTest(target=target):
                if self.report.exists():
                    import shutil
                    shutil.rmtree(self.report)
                self.plan()
                paths = {"source":self.source, "template":self.template,
                         "doc":self.baseline / "A.docx", "manifest":self.baseline / "revision.json",
                         "plan":self.report / "plan.json"}
                path = paths[target]
                saved = path.read_bytes()
                if target == "plan":
                    data = json.loads(saved); data["counts"] = {"unchanged":99}
                    path.write_bytes(revisions.encode(data))
                else:
                    path.write_bytes(saved + b"\n")
                with self.assertRaises(ValueError):
                    revisions.apply(self.report / "plan.json", self.next)
                self.assertFalse(self.next.exists())
                path.write_bytes(saved)

    def test_change_during_build_rolls_back_and_retry_succeeds(self):
        self.source.write_text("record_id,date,note\nA,2030-06-22,internal\nB,2030-07-01,draft\n")
        self.plan()
        original_write = revisions.write_document
        def mutate(text, path):
            original_write(text, path)
            self.source.write_text(self.source.read_text() + "C,2030-08-01,new\n")
        with patch("revisions.write_document", side_effect=mutate):
            with self.assertRaisesRegex(ValueError, "changed during generation"):
                revisions.apply(self.report / "plan.json", self.next)
        self.assertFalse(self.next.exists())
        fresh = self.root / "fresh-review"
        revisions.review(self.baseline, self.source, self.template, fresh)
        revisions.apply(fresh / "plan.json", self.next)
        self.assertEqual(len(list(self.next.glob("*.docx"))), 3)

    def test_generation_failure_no_partial_revision_or_overwrite(self):
        self.template.write_text("Changed title\n\n{{date}}")
        self.plan()
        with patch("revisions.write_document", side_effect=OSError("simulated disk failure")):
            with self.assertRaises(OSError):
                revisions.apply(self.report / "plan.json", self.next)
        self.assertFalse(self.next.exists())
        self.next.mkdir(); sentinel = self.next / "keep.txt"; sentinel.write_text("keep")
        with self.assertRaisesRegex(ValueError, "already exists"):
            revisions.apply(self.report / "plan.json", self.next)
        self.assertEqual(sentinel.read_text(), "keep")

    def test_missing_file_and_case_only_identifier_rename(self):
        (self.baseline / "A.docx").unlink()
        self.assertTrue(self.plan()["blocked"])
        self.source.write_text("record_id,date,note\na,2030-06-15,internal\nB,2030-07-01,draft\n")
        with self.assertRaisesRegex(ValueError, "letter case"):
            revisions.make_plan(self.baseline, self.source, self.template)

    def test_report_escapes_active_content_and_omits_absolute_paths(self):
        attack = '</script><img src=x onerror="alert(1)">'
        self.template.write_text("Service renewal\n\n" + attack + " {{date}}")
        plan = self.plan()
        page = render_page(plan)
        self.assertNotIn(attack, page)
        self.assertNotIn(str(self.root), page)
        self.assertIn("&lt;img", page)
        self.assertIn("connect-src 'none'", page)


if __name__ == "__main__":
    unittest.main()
